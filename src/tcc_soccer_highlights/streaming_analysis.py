"""Download e inferência sobrepostos, XLSX no GitHub e descarte dos vídeos temporários."""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from dataclasses import dataclass
import hashlib
import json
import logging
import os
from pathlib import Path, PurePosixPath
import threading

from openpyxl import load_workbook

from tcc_soccer_highlights.xlsx_publisher import XlsxPublisher

LOG = logging.getLogger(__name__)
FILES = ("1_224p.mkv", "2_224p.mkv", "1_720p.mkv", "2_720p.mkv")


@dataclass(frozen=True)
class Job:
    game: str
    split: str
    filename: str

    @property
    def key(self):
        return f"{self.game}/{self.filename}"

    def video(self, root):
        return root / self.game / self.filename

    def workbook(self, root):
        return root / self.game / (Path(self.filename).stem + "_analysis.xlsx")


def jobs_for(splits, files, games=None):
    from SoccerNet.utils import getListGames
    selected = set(games) if games else None
    result, seen = [], set()
    for split in splits:
        for game in getListGames(split=split):
            game = game.replace("\\", "/")
            if selected is not None and game not in selected:
                continue
            parts = PurePosixPath(game).parts
            if len(parts) != 3 or any(p in (".", "..") or ":" in p or "\\" in p for p in parts):
                raise ValueError(f"Identificador de partida inválido: {game}")
            for filename in files:
                job = Job(game, split, filename)
                if job.key not in seen:
                    seen.add(job.key)
                    result.append(job)
    if selected and selected - {job.game for job in result}:
        raise ValueError("Há partidas solicitadas que não pertencem aos conjuntos selecionados.")
    return result


def complete_workbook(path, args):
    """Não reutiliza arquivo quebrado, análise parcial ou planilha sem YOLO/cortes."""
    if not path.is_file():
        return False
    try:
        with path.open("rb") as source:
            workbook = load_workbook(source, read_only=True, data_only=True)
            try:
                if not {"metadata", "samples", "scenes", "events", "event_windows", "scene_summary"}.issubset(workbook.sheetnames):
                    return False
                rows = workbook["metadata"].iter_rows(values_only=True)
                meta = dict(zip(next(rows), next(rows)))
                duration = float(meta["duration_sec"])
                analyzed = float(meta.get("analyzed_duration_sec") or duration)
                expected = min(duration, args.max_seconds) if args.max_seconds else duration
                return (meta.get("status") == "ok" and duration > 0 and analyzed >= expected
                        and bool(meta.get("transnet_executed"))
                        and int(meta.get("yolo_stride", 0)) == args.stride
                        and Path(str(meta.get("yolo_model", ""))).name == args.model.name
                        and workbook["samples"].max_row > 1 and workbook["scenes"].max_row > 1)
            finally:
                workbook.close()
    except Exception as exc:
        # Inclui ZIP/XML inválido: nunca tratar XLSX incompleto como trabalho concluído.
        LOG.debug("Planilha não reutilizável %s: %s", path, exc)
        return False


@contextmanager
def pipeline_lock(state_dir):
    state_dir.mkdir(parents=True, exist_ok=True)
    with (state_dir / "run.lock").open("a+b") as handle:
        if os.name == "nt":
            import msvcrt
            acquire = lambda: msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            release = lambda: msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl
            acquire = lambda: fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
            release = lambda: fcntl.flock(handle, fcntl.LOCK_UN)
        try:
            if not (state_dir / "run.lock").stat().st_size:
                handle.write(b"0")
                handle.flush()
            handle.seek(0)
            acquire()
        except OSError as exc:
            raise RuntimeError("Já existe um pipeline usando este diretório de estado.") from exc
        try:
            yield
        finally:
            handle.seek(0)
            release()


class Journal:
    def __init__(self, root):
        self.root = root
        root.mkdir(parents=True, exist_ok=True)

    def path(self, job):
        return self.root / (hashlib.sha256(job.key.encode()).hexdigest() + ".json")

    def read(self, job):
        path = self.path(job)
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}

    def write(self, job, stage, **extra):
        path = self.path(job)
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(dict(video=job.key, stage=stage, **extra), ensure_ascii=False), encoding="utf-8")
        temporary.replace(path)


def run_pipeline(jobs, *, videos_root, output_root, state_dir, download, analyze, publish, valid, stop):
    """Um download antecipado por vez: no máximo dois vídeos no disco."""
    journal = Journal(state_dir / "jobs")
    pending, skipped = [], 0
    for job in jobs:
        output = job.workbook(output_root)
        previous = journal.read(job)
        if valid(output):
            if previous.get("stage") == "analyzed":
                LOG.info("Retomada: reenviando XLSX pronto, sem repetir inferência: %s", job.key)
                commit = publish(output)
                journal.write(job, "published", commit=commit)
                LOG.info("Retomada: XLSX salvo no GitHub (%s)", commit[:12])
            remove_video(job.video(videos_root), videos_root)
            skipped += 1
        else:
            pending.append(job)
    LOG.info("Planejamento: %d vídeos; %d já concluídos; %d pendentes", len(jobs), skipped, len(pending))
    if not pending:
        return 0
    jobs_iter = iter(pending)
    future = None
    executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix="download")
    try:
        current = next(jobs_iter)
        future = executor.submit(download, current)
        for index in range(len(pending)):
            if stop.is_set():
                break
            video = future.result()
            upcoming = next(jobs_iter, None)
            future = executor.submit(download, upcoming) if upcoming else None
            LOG.info("[%d/%d] Inferência: %s", index + 1, len(pending), current.key)
            # Se a escrita terminou antes de uma interrupção, a retomada só publica.
            journal.write(current, "analyzed")
            output = analyze(current, video)
            if not valid(output):
                raise RuntimeError(f"A inferência não produziu um XLSX completo: {output}")
            commit = publish(output)
            journal.write(current, "published", commit=commit)
            remove_video(video, videos_root)
            LOG.info("Salvo no GitHub (%s); vídeo removido: %s", commit[:12], current.key)
            current = upcoming
        return 0
    finally:
        stop.set()
        if future:
            future.cancel()
        executor.shutdown(wait=True, cancel_futures=True)
        # Só remove caminhos de vídeo criados e controlados por este pipeline.
        for job in pending:
            remove_video(job.video(videos_root), videos_root)
            remove_video(job.video(videos_root).with_suffix(".mkv.partial"), videos_root)


def remove_video(path, videos_root):
    resolved = path.resolve()
    resolved.relative_to(videos_root.resolve())
    if resolved.suffix.lower() not in (".mkv", ".partial"):
        raise ValueError("O descarte automático só remove vídeos temporários.")
    resolved.unlink(missing_ok=True)


def download_directory(project, state, requested=None):
    """Preferência local permite usar outro disco sem mover resultados ou checkpoints."""
    settings = state / "settings.json"
    saved = json.loads(settings.read_text(encoding="utf-8")) if settings.exists() else {}
    path = requested if requested is not None else saved.get("download_root", state / "videos")
    root = (project / Path(path)).resolve()
    if root == root.parent or root in (project.resolve(), state.resolve(), (project / "analysis_output").resolve()):
        raise ValueError("Escolha uma pasta exclusiva para os downloads temporários, não a raiz do disco ou do projeto.")
    return root


def prepare_download_directory(root, project, state, remember=False):
    marker = root / ".tcc_pipeline_owner.json"
    if marker.exists():
        owner = json.loads(marker.read_text(encoding="utf-8"))
        if owner.get("project") != str(project.resolve()):
            raise ValueError("A pasta de downloads já pertence ao pipeline de outro projeto.")
    elif root.exists() and any(root.iterdir()) and root != (state / "videos").resolve():
        raise ValueError("A pasta escolhida não está vazia. Use uma pasta exclusiva para os downloads do pipeline.")
    root.mkdir(parents=True, exist_ok=True)
    marker.write_text(json.dumps({"project": str(project.resolve())}), encoding="utf-8")
    if remember:
        settings = state / "settings.json"
        saved = json.loads(settings.read_text(encoding="utf-8")) if settings.exists() else {}
        saved["download_root"] = str(root)
        temporary = settings.with_suffix(".tmp")
        temporary.write_text(json.dumps(saved, ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(settings)


HF_VIDEO_REPO = "SoccerNet/SoccerNet_raw_HQ"


def hf_url(job, labels=False):
    from huggingface_hub import hf_hub_url
    return hf_hub_url("SoccerNet/SN-Labels" if labels else HF_VIDEO_REPO,
                      f"{job.game}/Labels-v2.json" if labels else job.key,
                      repo_type="dataset", revision="main" if labels else
                      ("videos-224p" if "224p" in job.filename else "videos-720p"))


def check_video_access(job, token):
    from huggingface_hub import get_hf_file_metadata
    if not token:
        raise PermissionError("Faça login local com `hf auth login` usando uma conta com acesso aprovado a "
                              "https://huggingface.co/datasets/SoccerNet/SoccerNet_raw_HQ .")
    try:
        get_hf_file_metadata(hf_url(job), token=token)
    except Exception as exc:
        raise PermissionError("Não foi possível acessar o vídeo no Hugging Face. Verifique o login, "
                              "a conexão e a aprovação de acesso ao dataset SoccerNet_raw_HQ.") from exc


def fetch_temporary(url, target, token, stop):
    """Baixa diretamente ao temporário próprio, sem reter o vídeo no cache do Hub."""
    import requests
    target.parent.mkdir(parents=True, exist_ok=True)
    temporary = target.with_suffix(target.suffix + ".partial")
    headers = {"Accept-Encoding": "identity"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    try:
        # Requests remove Authorization em redirecionamentos para outro host.
        with requests.get(url, headers=headers, stream=True, timeout=(20, 60)) as response:
            if response.status_code in (401, 403):
                raise PermissionError("Hugging Face recusou o download; confira o login e o acesso ao dataset.")
            response.raise_for_status()
            length = int(response.headers.get("Content-Length", "0"))
            received = 0
            with temporary.open("wb") as handle:
                for chunk in response.iter_content(chunk_size=1024 * 1024):
                    if stop.is_set():
                        raise RuntimeError("Download interrompido.")
                    handle.write(chunk)
                    received += len(chunk)
            if not received or (length and received != length):
                raise OSError("Download vazio ou incompleto.")
        temporary.replace(target)
    finally:
        temporary.unlink(missing_ok=True)


def downloader_for(videos_root, token, stop):
    def download(job):
        LOG.info("Download: %s", job.key)
        label_path = videos_root / job.game / "Labels-v2.json"
        if not label_path.exists():
            fetch_temporary(hf_url(job, labels=True), label_path, token, stop)
        labels = json.loads(label_path.read_text(encoding="utf-8"))
        if not isinstance(labels.get("annotations"), list):
            raise ValueError("As anotações SoccerNet estão ausentes ou inválidas.")
        fetch_temporary(hf_url(job), job.video(videos_root), token, stop)
        LOG.info("Download concluído: %s", job.key)
        return job.video(videos_root)

    return download


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=Path.cwd())
    parser.add_argument("--state-dir", type=Path, default=Path(".pipeline_state"))
    parser.add_argument("--download-root", type=Path,
                        help="Pasta exclusiva para vídeos temporários; a escolha fica salva localmente")
    parser.add_argument("--model", type=Path, default=Path("best.pt"))
    parser.add_argument("--files", nargs="+", choices=FILES, default=["1_224p.mkv", "2_224p.mkv"])
    parser.add_argument("--splits", nargs="+", choices=["train", "valid", "test"], default=["train", "valid", "test"])
    parser.add_argument("--games", nargs="+", help="Identificadores completos liga/temporada/partida")
    parser.add_argument("--branch", default="main")
    parser.add_argument("--remote", default="origin")
    parser.add_argument("--stride", type=int, default=5)
    parser.add_argument("--batch-size", type=int, default=16)
    parser.add_argument("--event-window", type=float, default=10.)
    parser.add_argument("--transnet-threshold", type=float, default=.5)
    parser.add_argument("--device", choices=["auto", "cuda", "cpu"], default="auto")
    parser.add_argument("--limit", type=int)
    parser.add_argument("--max-seconds", type=float, help="Teste parcial, salvo em analysis_output/pipeline_tests/")
    parser.add_argument("--dry-run", action="store_true", help="Somente listar trabalhos pendentes")
    return parser.parse_args(argv)


def inference_device(requested):
    import torch
    if requested == "cpu":
        return "cpu"
    try:
        if not torch.cuda.is_available():
            raise RuntimeError("CUDA indisponível")
        # Detectar CUDA não garante que os kernels instalados suportem esta GPU.
        torch.ones(1, device="cuda").sum().item()
        return "cuda"
    except RuntimeError as exc:
        if requested == "cuda":
            raise RuntimeError("A instalação atual do PyTorch não consegue executar nesta GPU; use --device cpu.") from exc
        LOG.warning("CUDA não executável nesta instalação; usando CPU para a inferência.")
        return "cpu"


def main(argv=None):
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    args = parse_args(argv)
    project = args.project_root.resolve()
    args.model = (project / args.model).resolve()
    state = (project / args.state_dir).resolve()
    state.relative_to(project)
    if state == project or state.name == ".git" or not state.name.startswith("."):
        raise ValueError("O estado deve ficar em um subdiretório oculto próprio do projeto.")
    args.videos_root = download_directory(project, state, args.download_root)
    args.output_root = project / "analysis_output"
    if args.max_seconds:
        args.output_root /= "pipeline_tests"
    args.skip_transnet = False
    if args.stride < 1 or args.batch_size < 1 or (args.limit is not None and args.limit < 1):
        raise ValueError("Stride, batch-size e limit devem ser positivos.")
    if args.max_seconds is not None and args.max_seconds <= 0:
        raise ValueError("max-seconds deve ser positivo.")
    if not args.model.is_file():
        raise FileNotFoundError(f"Modelo não encontrado: {args.model}")
    jobs = jobs_for(args.splits, args.files, args.games)
    if args.limit:
        jobs = jobs[:args.limit]
    valid = lambda path: complete_workbook(path, args)
    if args.dry_run:
        pending = [j for j in jobs if not valid(j.workbook(args.output_root))]
        LOG.info("%d vídeos selecionados; %d completos localmente; %d a baixar", len(jobs), len(jobs)-len(pending), len(pending))
        LOG.info("Downloads temporários: %s", args.videos_root)
        for job in pending[:10]:
            LOG.info("Pendente: %s", job.key)
        return 0
    with pipeline_lock(state):
        prepare_download_directory(args.videos_root, project, state, remember=args.download_root is not None)
        LOG.info("Downloads temporários: %s", args.videos_root)
        publisher = XlsxPublisher(project, state, args.remote, args.branch)
        publisher.preflight()
        from huggingface_hub import get_token
        token = get_token()
        first_pending = next((job for job in jobs if not valid(job.workbook(args.output_root))), None)
        if first_pending:
            check_video_access(first_pending, token)
        stop = threading.Event()
        stop_file = state / "STOP"
        stop_file.unlink(missing_ok=True)

        def watch_stop():
            while not stop.wait(1):
                if stop_file.exists():
                    LOG.info("Parada solicitada: salvando a inferência atual antes de encerrar.")
                    stop.set()

        threading.Thread(target=watch_stop, name="stop-signal", daemon=True).start()
        download = downloader_for(args.videos_root, token, stop)
        from tcc_soccer_highlights.batch_analysis import VideoInfo, analyze_video, load_yolo
        device = inference_device(args.device)
        model = load_yolo(args.model, .25, .45, device)
        LOG.info("Inferência em %s; publicação em %s/%s; no máximo 2 vídeos temporários", device, args.remote, args.branch)

        def analyze(job, video):
            info = VideoInfo(video, video.parent, int(job.filename[0]), f"{job.game}/{video.stem}")
            return analyze_video(info, args, model, device)

        return run_pipeline(jobs, videos_root=args.videos_root, output_root=args.output_root,
                            state_dir=state, download=download, analyze=analyze, publish=publisher.publish,
                            valid=valid, stop=stop)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        LOG.warning("Execução interrompida. Os XLSX concluídos serão reutilizados na retomada.")
        raise SystemExit(130)
    except Exception:
        LOG.exception("Pipeline interrompido; os XLSX locais foram preservados para retomada.")
        raise SystemExit(1)
